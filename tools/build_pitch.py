"""Build matching editable PPTX and PDF slides for the six-minute pitch."""
from pathlib import Path
from xml.sax.saxutils import escape
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts' / 'presentacion'
OUT.mkdir(parents=True, exist_ok=True)
W, H = 960, 540
BG, DARK, GREEN, MUTED, GOLD, WHITE = 'F5F5EF', '163D34', '246D57', '53665D', 'D5AD62', 'FFFFFF'
pdfmetrics.registerFont(TTFont('Arial', 'C:/Windows/Fonts/arial.ttf'))
pdfmetrics.registerFont(TTFont('ArialBold', 'C:/Windows/Fonts/arialbd.ttf'))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
pdf = canvas.Canvas(str(OUT / 'SpendWiseAI_Pitch_6min.pdf'), pagesize=(W, H))
pdf.setTitle('SpendWise AI | Pitch de seis minutos')
pdf.setAuthor('Equipo SpendWise AI')
slide = None


def rect(x, y, w, h, color, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    if radius:
        shape.adjustments[0] = .08
    pdf.setFillColor(HexColor('#'+color))
    if radius: pdf.roundRect(x, H-y-h, w, h, 9, fill=1, stroke=0)
    else: pdf.rect(x, H-y-h, w, h, fill=1, stroke=0)


def text(x, y, w, value, size=20, color=DARK, bold=False):
    # Explicit wrapping shared by both formats keeps the layouts consistent.
    font = 'ArialBold' if bold else 'Arial'
    lines = []
    for paragraph in value.split('\n'):
        current = ''
        for word in paragraph.split():
            test = (current+' '+word).strip()
            if pdfmetrics.stringWidth(test,font,size) > w and current:
                lines.append(current); current=word
            else: current=test
        lines.append(current)
    height = len(lines)*size*1.25+6
    box=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w+4),Pt(height))
    tf=box.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.word_wrap=False
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold
        p.font.color.rgb=RGBColor.from_string(color)
        p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(size*1.25)
        pdf.setFillColor(HexColor('#'+color));pdf.setFont(font,size)
        pdf.drawString(x,H-y-size-i*size*1.25,line)
    assert y+height < H, (value,y,height)


def begin(number, label, title, timing, notes, dark=False):
    global slide
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    rect(0,0,W,H,DARK if dark else BG)
    text(48,28,650,'SPENDWISE AI  /  '+label,11,GOLD if dark else GREEN,True)
    text(48,66,865,title,34,WHITE if dark else DARK,True)
    rect(48,490,864,1,GREEN if dark else 'D7DED5')
    text(48,506,670,'PROTOTIPO ACADÉMICO  ·  '+timing,10,'C7D8CC' if dark else MUTED)
    text(850,502,65,f'{number:02d} / 07',12,GOLD if dark else GREEN,True)
    slide.notes_slide.notes_text_frame.text=timing+'\n\n'+notes


def end(): pdf.showPage()


notes=[
'''¿Alguna vez han llegado al final del mes sin entender en qué se les fue la plata?
Muchas personas conocen cuánto reciben, pero tienen sus gastos repartidos entre recuerdos, notas y movimientos bancarios. Organizar esa información requiere tiempo y, cuando finalmente lo hacen, todavía queda una pregunta: ¿qué puedo cambiar para mejorar?
De esa necesidad nace SpendWise AI, una aplicación web que permite describir los gastos en lenguaje cotidiano, convertirlos en un presupuesto revisable y conversar sobre posibles ajustes.
Nuestro propósito es ayudar al usuario a pasar de registrar su dinero a entenderlo y tomar decisiones.''',
'''SpendWise organiza la experiencia en cuatro secciones: Registrar, Historial, Comparar y Planear.
En Registrar, el usuario escribe sus ingresos y gastos como los contaría normalmente. Gemini interpreta el texto y propone una organización de los movimientos.
Antes de calcular, presentamos una pantalla de revisión. Allí se pueden corregir montos, cambiar categorías o excluir gastos. Después de confirmar, se obtiene el presupuesto y se puede guardar.
El Historial permite consultar meses anteriores, mientras que Comparar ayuda a entender sus diferencias.
Elegimos una aplicación web porque facilita el acceso desde el navegador y reúne todo el recorrido en una interfaz que también puede utilizarse desde el celular.''',
'''[Mostrar un presupuesto preparado y abrir Planear.]
Veamos cómo se utiliza. Aquí tenemos un presupuesto que el usuario ya revisó y guardó.
En Planear podemos seleccionarlo y escribir: «Quiero ahorrar más, pero no puedo reducir el arriendo ni el transporte. Propón una meta realista».
El asistente recibe los movimientos del mes y un resumen calculado por la aplicación. Con esa información puede explicar alternativas y proponer ajustes concretos.
La conversación continúa: si respondemos «Eso me parece demasiado, empecemos con algo más pequeño», conserva el contexto y puede adaptar la propuesta.
Al revisar una simulación, el sistema muestra el ahorro propuesto y su efecto sobre el saldo.
Este es el flujo AI-native: el lenguaje natural es parte central de la interacción. La IA interpreta y propone, el código valida y calcula, y el usuario conserva la decisión final.
[Si tarda, mostrar una conversación obtenida previamente y aclarar que es una ejecución anterior.]''',
'''La arquitectura separa la interfaz, la lógica de la aplicación y los servicios externos.
La interfaz utiliza HTML, CSS y JavaScript. El servidor está construido con Python y FastAPI, y coordina las solicitudes, las validaciones y la comunicación con Gemini. Los presupuestos se almacenan en SQLite.
La clave de Gemini permanece en el servidor.
Una decisión central fue separar las propuestas de los cálculos. Por ejemplo, si Gemini propone reducir un gasto, Python verifica que ese movimiento exista y que el recorte no supere su valor original.
Los cálculos de la simulación utilizan Decimal para trabajar con importes monetarios. Además, conversar o revisar una propuesta no modifica automáticamente el presupuesto guardado.''',
'''Evaluar un sistema con IA requiere comprobar dos cosas: que el software funcione y que las respuestas del modelo sean útiles y correctas.
En extracción, los criterios incluyen reconocer los montos, conservar los movimientos mencionados y manejar información incompleta.
En conversación, nos interesa comprobar que el asistente use el presupuesto, considere las restricciones del usuario y mantenga el contexto.
Para la nueva sección de planificación ejecutamos seis pruebas automatizadas. Verifican los cálculos, el rechazo de recortes inválidos, el envío del contexto, el consentimiento y el manejo de errores. También comprobamos el recorrido en navegador y la presentación móvil.
Estas pruebas utilizan respuestas simuladas: verifican nuestra implementación, pero no certifican la calidad de Gemini real.
Como gates técnicos propuestos, tenemos consentimiento, validación de propuestas y revisión humana. La evaluación repetida con Gemini, las pruebas con usuarios y el contraste con la rúbrica oficial siguen siendo tareas pendientes.''',
'''Una dificultad importante fue manejar los tiempos de espera y los errores del proveedor. Encontramos solicitudes que agotaban el tiempo disponible y respuestas de error del servidor.
Aprendimos que debemos distinguir entre tener una clave configurada y recibir una respuesta válida. También que los mensajes de error deben ayudar a entender el problema sin exponer información sensible.
Otra dificultad fue la interacción en Planear. Inicialmente era un formulario que interpretaba solicitudes aisladas. Eso limitaba la posibilidad de corregir una propuesta.
Lo transformamos en una conversación para que el usuario pudiera expresar preferencias y pedir alternativas.
El aprendizaje principal fue que construir una aplicación con IA exige diseñar también la revisión, las correcciones y el comportamiento cuando la IA falla.''',
'''El valor de SpendWise está en conectar tres pasos: contar cómo usamos el dinero, entender el presupuesto y explorar qué podríamos cambiar.
El prototipo tiene límites claros. Depende de la disponibilidad de Gemini; la conversación permanece mientras la página esté abierta, y las estimaciones de eventos deben confirmarse porque no son precios consultados en tiempo real.
El siguiente paso es probarlo con usuarios y medir cuántas correcciones necesitan, cuánto tardan en completar el recorrido y si las propuestas les ayudan a decidir.
No queremos que el usuario acepte una respuesta automáticamente. Queremos que pueda entenderla, cuestionarla y ajustarla a su realidad.
SpendWise AI: entiende tu dinero, explora tus opciones y decide con más claridad.'''
]
times=['0:00–0:45 · 45 s','0:45–1:30 · 45 s','1:30–2:30 · 60 s','2:30–3:20 · 50 s','3:20–4:20 · 60 s','4:20–5:10 · 50 s','5:10–6:00 · 50 s']

begin(1,'EL PROBLEMA','¿En qué se me fue la plata?',times[0],notes[0],True)
text(48,156,560,'De notas sueltas a decisiones\ncon más claridad.',29,WHITE)
rect(655,152,257,256,GREEN,True)
text(678,178,211,'GASTOS DISPERSOS',13,GOLD,True)
text(678,218,210,'Recuerdos\nNotas\nMovimientos bancarios',22,WHITE)
text(48,270,535,'SpendWise AI',46,GOLD,True)
text(48,341,530,'Describe tus gastos. Revisa tu presupuesto.\nConversa sobre qué puedes mejorar.',21,WHITE)
end()

begin(2,'LA SOLUCIÓN','Un recorrido para entender tu dinero',times[1],notes[1])
for i,(title,desc) in enumerate([('Registrar','Texto cotidiano → revisión → presupuesto'),('Historial','Consulta los meses que ya guardaste.'),('Comparar','Entiende cómo cambiaron tus gastos.'),('Planear','Conversa y explora ajustes posibles.')]):
    x=48+(i%2)*442;y=151+(i//2)*131
    rect(x,y,422,114,WHITE,True);text(x+20,y+13,380,f'0{i+1}  {title}',23,GREEN,True);text(x+20,y+52,372,desc,18)
text(48,432,864,'WEB APP  ·  Acceso desde el navegador y diseño adaptable al celular.',17,MUTED)
end()

begin(3,'DEMO / AI-NATIVE','Una conversación que lleva a un plan',times[2],notes[2])
rect(48,144,565,101,GREEN,True)
text(68,157,524,'TÚ',11,'E4D4A9',True)
text(68,180,518,'Quiero ahorrar más, pero no puedo reducir el arriendo ni el transporte.',21,WHITE)
rect(48,259,565,94,WHITE,True)
text(68,274,522,'CONTEXTO PARA GEMINI',11,GREEN,True)
text(68,297,520,'Presupuesto seleccionado + conversación + restricciones del usuario.',20)
rect(48,367,565,70,'E6ECE2',True)
text(68,381,520,'«Empecemos con una meta más pequeña».',21,GREEN,True)
for i,(a,b) in enumerate([('IA','Interpreta y propone'),('CÓDIGO','Valida y calcula'),('PERSONA','Revisa y decide')]):
    y=151+i*98;text(651,y,250,a,12,GREEN,True);text(651,y+25,260,b,22)
text(651,449,255,'Propuestas sin guardado automático.',13,MUTED)
end()

begin(4,'ARQUITECTURA','Responsabilidades separadas',times[3],notes[3])
for x,title,sub in [(48,'Navegador','HTML · CSS · JavaScript'),(352,'Servidor','Python · FastAPI'),(656,'Gemini','Interpretación y diálogo')]:
    rect(x,155,256,102,WHITE,True);text(x+18,172,221,title,25,GREEN,True);text(x+18,212,221,sub,15)
text(315,184,32,'↔',25,GREEN);text(618,184,32,'↔',25,GREEN)
rect(352,296,256,90,'E6ECE2',True);text(370,307,220,'SQLite',23,GREEN,True);text(370,345,220,'Presupuestos y movimientos',14)
text(461,263,50,'↕',25,GREEN)
text(48,296,261,'Clave en el servidor.\nImportes con Decimal.',20)
text(656,292,252,'Recortes validados contra gastos existentes.',20)
rect(48,415,864,49,DARK,True);text(67,425,825,'La propuesta de la IA pasa por validación antes de mostrar su efecto.',19,WHITE)
end()

begin(5,'EVALS Y GATES','Evidencia actual y validación pendiente',times[4],notes[4])
rect(48,144,420,247,WHITE,True);text(69,162,374,'VERIFICADO EN PLANIFICACIÓN',13,GREEN,True)
text(69,196,80,'6',62,GREEN,True);text(158,210,270,'pruebas automatizadas\ncon respuestas simuladas',19)
text(69,285,376,'Cálculos · recortes inválidos · contexto\nConsentimiento · errores · sin escrituras',17)
text(69,348,375,'Recorrido en navegador y vista móvil.',17,MUTED)
rect(489,144,423,247,'EAECE0',True);text(510,162,381,'PENDIENTE DE EVIDENCIA',13,GREEN,True)
text(510,204,368,'Calidad de Gemini en pruebas repetidas.\nValidación con usuarios reales.\nContraste con la rúbrica oficial.',21)
rect(48,414,864,53,DARK,True);text(68,424,822,'Gates propuestos: consentimiento → validación → revisión humana.',19,WHITE)
end()

begin(6,'APRENDIZAJES','Diseñar también cuando la IA falla',times[5],notes[5])
for i,(a,b) in enumerate([('Timeouts y errores del proveedor','Distinguir conexión, disponibilidad y respuesta válida.'),('Solicitudes aisladas en Planear','Conversación con contexto y ajustes de seguimiento.'),('Propuestas difíciles de comprobar','Validaciones y simulaciones con cálculos verificables.')]):
    y=151+i*96
    rect(48,y,864,80,WHITE,True)
    text(67,y+13,340,a,20,GREEN,True);text(438,y+13,447,b,19)
end()

begin(7,'CIERRE','Entiende. Explora. Decide.',times[6],notes[6],True)
text(48,151,851,'Una conversación para tomar decisiones\ncon más claridad sobre tu dinero.',31,WHITE)
rect(48,260,420,171,GREEN,True);text(69,278,370,'LÍMITES DEL PROTOTIPO',13,GOLD,True)
text(69,314,366,'Disponibilidad de Gemini.\nMemoria mientras la página está abierta.\nEstimaciones por confirmar.',18,WHITE)
rect(489,260,423,171,GREEN,True);text(510,278,376,'SIGUIENTE PASO',13,GOLD,True)
text(510,314,375,'Probar con usuarios y medir:\ncorrecciones, tiempo del recorrido\ny utilidad de las propuestas.',19,WHITE)
text(48,451,863,'SPENDWISE AI  ·  Tú conservas la decisión final.',20,GOLD,True)
end()

prs.core_properties.title='SpendWise AI — Pitch de seis minutos'
prs.core_properties.subject='Prototipo académico: flujo AI-native, arquitectura, evals y gates'
prs.save(OUT / 'SpendWiseAI_Pitch_6min.pptx');pdf.save()
md='# SpendWise AI · Guion de seis minutos\n\nLos bloques suman 6:00. Ensayar con cronómetro; el tiempo incluye una demo breve.\n\n'
for i,note in enumerate(notes): md+=f'## Diapositiva {i+1} · {times[i]}\n\n{note}\n\n'
(OUT/'Guion_6_minutos.md').write_text(md,encoding='utf-8')
print('Generados:',*[str(p) for p in OUT.iterdir()],sep='\n')
